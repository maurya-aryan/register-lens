"""Plain-code logic: drug matching, row validation, DVDMS reconciliation, export, alerts.

Gemini reads the page; everything numeric after that is deterministic code so the
numbers are auditable.
"""
import io
import json
import re
from datetime import date, datetime
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rapidfuzz import fuzz, process
from sqlmodel import Session, select

from . import config
from .db import PHC, DvdmsStock, Medicine
from .gemini_reader import generate_json_text

MATCH_THRESHOLD = 80


# ---------- drug matching ----------
def _variants(m: Medicine) -> List[str]:
    out = [m.name, f"{m.name} {m.strength}"]
    out += json.loads(m.aliases)
    return out


def build_index(meds: List[Medicine]):
    choices, owner = [], []
    for m in meds:
        for v in _variants(m):
            choices.append(v.lower())
            owner.append(m.code)
    return choices, owner


def match_drug(raw: str, strength: Optional[str], index) -> tuple:
    choices, owner = index
    q = (raw or "").strip().lower()
    if strength:
        q_full = f"{q} {strength.lower()}"
    else:
        q_full = q
    if not q:
        return None, 0
    best = process.extractOne(q_full, choices, scorer=fuzz.WRatio)
    if best and best[1] >= MATCH_THRESHOLD:
        return owner[best[2]], int(best[1])
    return None, int(best[1]) if best else 0


def gemini_match_names(raws: List[str], meds: List[Medicine]) -> Dict[str, Optional[str]]:
    """Fallback: ask Gemini to map hard handwritten names onto our catalogue codes."""
    if not raws:
        return {}
    catalogue = "\n".join(f"{m.code}: {m.name} {m.strength} {m.form}" for m in meds)
    prompt = (
        "Map each handwritten medicine name to the best matching catalogue code, or null if none fits.\n"
        "Names may be Hindi, English, abbreviations or brand names. Do not force a match.\n\n"
        f"CATALOGUE:\n{catalogue}\n\nNAMES:\n" + "\n".join(f"- {r}" for r in raws) +
        "\n\nReturn JSON object: {\"<name>\": \"<code or null>\"}"
    )
    data, _ = generate_json_text(prompt, 0.0)
    if not isinstance(data, dict):
        return {}
    valid = {m.code for m in meds}
    return {k: (v if v in valid else None) for k, v in data.items()}


# ---------- validation ----------
def _parse_ym(s: Optional[str]) -> Optional[date]:
    if not s:
        return None
    m = re.match(r"^(\d{4})-(\d{1,2})$", s.strip())
    if not m:
        return None
    y, mo = int(m.group(1)), int(m.group(2))
    if not 1 <= mo <= 12:
        return None
    return date(y, mo, 28)


def validate_row(r: dict) -> List[dict]:
    flags = []
    conf = r["confidence"]
    for key, label in (("drug", "drug name"), ("batch", "batch"), ("expiry", "expiry"), ("qty", "quantities")):
        if conf.get(key, 1) < config.LOW_CONF:
            flags.append({"code": f"low_{key}", "level": "warn", "msg": f"Unsure about {label}. Please check."})
    o, rc, i, c = r.get("opening"), r.get("received"), r.get("issued"), r.get("closing")
    if None in (o, i, c):
        flags.append({"code": "missing_qty", "level": "warn", "msg": "A quantity could not be read."})
    else:
        expected = o + (rc or 0) - i
        if expected != c:
            flags.append({"code": "balance", "level": "error",
                          "msg": f"Balance does not add up: {o} + {rc or 0} - {i} = {expected}, but closing is {c}."})
        if i > o + (rc or 0):
            flags.append({"code": "over_issue", "level": "error", "msg": "Issued more than was in stock."})
    exp = _parse_ym(r.get("expiry"))
    if r.get("expiry") and exp is None:
        flags.append({"code": "bad_expiry", "level": "warn", "msg": "Expiry date is not a valid month."})
    elif exp and exp < date.today():
        flags.append({"code": "expired", "level": "error", "msg": "This batch is already expired."})
    elif exp and (exp - date.today()).days <= 60:
        flags.append({"code": "expiring", "level": "info", "msg": "Expires within 60 days."})
    return flags


def prepare_rows(page_rows: List[dict], meds: List[Medicine], dv_batches: Optional[Dict[str, List[tuple]]] = None) -> List[dict]:
    index = build_index(meds)
    by_code = {m.code: m for m in meds}
    out, pending = [], []
    for i, r in enumerate(page_rows):
        code, score = match_drug(r["drug_name_raw"], r.get("strength_raw"), index)
        row = {**r, "id": i, "med_code": code, "match_score": score, "matched_by": "fuzzy" if code else None}
        out.append(row)
        if code is None:
            pending.append(row)
    if pending:
        mapping = gemini_match_names([p["drug_name_raw"] for p in pending], meds)
        for p in pending:
            code = mapping.get(p["drug_name_raw"])
            if code:
                p["med_code"], p["matched_by"] = code, "gemini"
    for row in out:
        m = by_code.get(row["med_code"]) if row["med_code"] else None
        row["med_name"] = f"{m.name} {m.strength}" if m else None
        row["unit"] = m.unit if m else None
        row["flags"] = validate_row(row)
        row["flags"] += batch_crosscheck(row, (dv_batches or {}).get(row["med_code"] or "", []))
        if m is None:
            row["flags"].append({"code": "unmatched", "level": "warn", "msg": "Could not match this to a known medicine."})
    return out


def batch_crosscheck(row: dict, known: List[tuple]) -> List[dict]:
    """Compare the register's batch/expiry with the batches DVDMS already lists for this centre.
    Catches single-character misreads that the model's own confidence misses."""
    b = (row.get("batch") or "").replace(" ", "").upper()
    if not b or not known:
        return []
    exact = [k for k in known if k[0].upper() == b]
    if exact:
        exp = row.get("expiry")
        if exp and exact[0][1] and exp != exact[0][1]:
            return [{"code": "expiry_mismatch", "level": "warn",
                     "msg": f"DVDMS lists expiry {exact[0][1]} for batch {b}, the register reads {exp}. Please check."}]
        return []
    near = [k for k in known if fuzz.ratio(k[0].upper(), b) >= 66]
    if near:
        return [{"code": "batch_near", "level": "warn",
                 "msg": f"Batch {b} looks like {near[0][0]} in DVDMS. Check for a misread digit."}]
    return [{"code": "new_batch", "level": "info", "msg": "New batch, not yet in DVDMS."}]


# ---------- reconciliation ----------
def dvdms_batches(session: Session, phc_id: str) -> Dict[str, List[tuple]]:
    out: Dict[str, List[tuple]] = {}
    for s in session.exec(select(DvdmsStock).where(DvdmsStock.phc_id == phc_id)).all():
        out.setdefault(s.med_code, []).append((s.batch, s.expiry))
    if out:
        return out
    fac = session.get(PHC, phc_id)
    if fac:
        from .stock import simulated
        for code, v in simulated(fac.id, fac.kind, fac.load).items():
            out[code] = [(b["batch"], b["expiry"][:7]) for b in v["batches"]]
    return out


def match_phc(session: Session, written: Optional[str]) -> Optional[str]:
    """Pick the health centre from the facility name written on the page."""
    if not written:
        return None
    phcs = session.exec(select(PHC).where(PHC.demo == True)).all()  # noqa: E712
    best = process.extractOne(written.lower(), [p.name.lower() for p in phcs], scorer=fuzz.WRatio)
    if best and best[1] >= 90:
        return phcs[best[2]].id
    return None


def reconcile(session: Session, phc_id: str, rows: List[dict]) -> dict:
    phc = session.get(PHC, phc_id)
    meds = {m.code: m for m in session.exec(select(Medicine)).all()}
    from .stock import facility_stock
    fstock = facility_stock(session, phc) if phc else {}
    items = []
    for r in rows:
        code = r.get("med_code")
        if not code or code not in meds or r.get("closing") is None:
            continue
        m = meds[code]
        stock = session.exec(select(DvdmsStock).where(DvdmsStock.phc_id == phc_id, DvdmsStock.med_code == code)).all()
        dvdms_qty = sum(s.qty for s in stock) if stock else int(fstock.get(code, {}).get("dvdms", 0))
        baseline = fstock.get(code, {}).get("daily") or max(1, round(m.base_daily_issue * (phc.load if phc else 1.0)))
        issued = r.get("issued") or 0
        daily = issued if issued > 0 else baseline
        closing = r["closing"]
        days_real = round(closing / daily, 1) if daily else None
        days_dvdms = round(dvdms_qty / daily, 1) if daily else None
        drift = dvdms_qty - closing
        drift_pct = round(100 * drift / dvdms_qty, 1) if dvdms_qty else 0.0
        exp = _parse_ym(r.get("expiry"))
        expiring = bool(exp and 0 <= (exp - date.today()).days <= 60)
        if days_real is not None and days_real <= 7:
            status = "critical"
        elif days_real is not None and days_real <= 15:
            status = "warning"
        else:
            status = "ok"
        items.append({
            "med_code": code, "med_name": f"{m.name} {m.strength}", "unit": m.unit,
            "batch": r.get("batch"), "expiry": r.get("expiry"),
            "register_closing": closing, "dvdms_qty": dvdms_qty, "drift": drift, "drift_pct": drift_pct,
            "daily_use": daily, "daily_use_source": "register" if issued > 0 else "baseline",
            "days_real": days_real, "days_dvdms": days_dvdms, "status": status,
            "expiring_soon": expiring,
            "phantom_stock": drift_pct >= 40,
        })
    order = {"critical": 0, "warning": 1, "ok": 2}
    items.sort(key=lambda x: (order[x["status"]], -x["drift_pct"]))
    summary = {
        "medicines": len(items),
        "critical": sum(1 for i in items if i["status"] == "critical"),
        "warning": sum(1 for i in items if i["status"] == "warning"),
        "phantom": sum(1 for i in items if i["phantom_stock"]),
        "expiring": sum(1 for i in items if i["expiring_soon"]),
        "units_overstated": sum(max(0, i["drift"]) for i in items),
    }
    return {"phc": {"id": phc.id, "name": phc.name, "district": phc.district, "state": phc.state, "kind": phc.kind} if phc else None,
            "items": items, "summary": summary}


# ---------- export ----------
def build_bulk_entry(phc: PHC, rows: List[dict], entry_date: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DVDMS Bulk Entry"
    head = ["Facility", "District", "Entry Date", "Drug Code", "Drug Name", "Strength/Form",
            "Batch No", "Expiry (MM/YYYY)", "Opening", "Received", "Issued", "Closing Stock", "Unit"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        if not r.get("med_code"):
            continue
        exp = r.get("expiry") or ""
        exp_fmt = f"{exp[5:7]}/{exp[:4]}" if re.match(r"^\d{4}-\d{2}$", exp) else ""
        ws.append([phc.name, phc.district, entry_date, r["med_code"], r.get("med_name"),
                   "", r.get("batch") or "", exp_fmt, r.get("opening"), r.get("received"),
                   r.get("issued"), r.get("closing"), r.get("unit")])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(34, max(len(str(c.value or "")) for c in col) + 2))
    note = wb.create_sheet("About")
    note.append(["Generated by Register Lens. DVDMS-compatible template (sample data). Review before uploading."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- alert ----------
def _fallback_alert(phc: dict, summary: dict, items: List[dict]) -> dict:
    top = [i for i in items if i["status"] == "critical"][:3]
    en_lines = [f"- {i['med_name']}: about {i['days_real']} days left (DVDMS shows {i['days_dvdms']})" for i in top]
    hi_lines = [f"- {i['med_name']}: लगभग {i['days_real']} दिन का स्टॉक बचा (DVDMS में {i['days_dvdms']} दिन दिख रहा है)" for i in top]
    en = f"Stock alert for {phc['name']}, {phc['district']}.\n" + ("\n".join(en_lines) or "No critical shortages today.") + \
         "\nPlease review and arrange supply."
    hi = f"{phc['name']}, {phc['district']} के लिए स्टॉक अलर्ट।\n" + ("\n".join(hi_lines) or "आज कोई गंभीर कमी नहीं है।") + \
         "\nकृपया समीक्षा करें और आपूर्ति की व्यवस्था करें।"
    return {"en": en, "hi": hi, "source": "template"}


LANGUAGES = {
    "hi": ("Hindi", "हिन्दी"), "en": ("English", "English"), "ur": ("Urdu", "اردو"), "bn": ("Bengali", "বাংলা"),
    "ta": ("Tamil", "தமிழ்"), "te": ("Telugu", "తెలుగు"), "mr": ("Marathi", "मराठी"), "gu": ("Gujarati", "ગુજરાતી"),
    "kn": ("Kannada", "ಕನ್ನಡ"), "ml": ("Malayalam", "മലയാളം"), "or": ("Odia", "ଓଡ଼ିଆ"), "pa": ("Punjabi", "ਪੰਜਾਬੀ"),
    "as": ("Assamese", "অসমীয়া"), "ne": ("Nepali", "नेपाली"), "bho": ("Bhojpuri", "भोजपुरी"), "awa": ("Awadhi", "अवधी"),
}
STATE_LANGS = {"Uttar Pradesh": ["hi", "ur", "en"]}

_ALERT_CACHE: dict = {}


def write_alert(rec: dict, langs: Optional[List[str]] = None) -> dict:
    """Alert text per language code. Numbers come only from the reconciliation."""
    phc, items, summary = rec["phc"], rec["items"], rec["summary"]
    langs = [l for l in (langs or ["hi", "en"]) if l in LANGUAGES] or ["hi", "en"]
    if not phc:
        return {"texts": {}, "source": "none"}
    ckey = json.dumps([phc["id"], langs, [(i["med_code"], i["register_closing"], i["dvdms_qty"], i["status"]) for i in items]])
    if ckey in _ALERT_CACHE:
        return _ALERT_CACHE[ckey]
    facts = [
        {k: i[k] for k in ("med_name", "register_closing", "dvdms_qty", "days_real", "days_dvdms", "status", "expiring_soon", "expiry")}
        for i in items if i["status"] != "ok" or i["expiring_soon"]
    ][:8]
    lang_list = ", ".join(f'"{c}" = {LANGUAGES[c][0]} ({LANGUAGES[c][1]} script)' for c in langs)
    prompt = (
        "You write short, clear stock alerts from a PHC pharmacist to the District Drug Store Officer in India.\n"
        f"Facility: {phc['name']}, {phc['district']}, {phc['state']}.\n"
        f"Facts (use ONLY these numbers, do not invent any): {json.dumps(facts, ensure_ascii=False)}\n"
        f"Write the same alert in each of these languages: {lang_list}. Use each language's own script; keep medicine "
        "names recognisable (you may keep them in English). Max 6 lines each. Start with the most urgent item. "
        "State days of stock left according to the register vs what DVDMS shows. End with one polite action request.\n"
        "Return a JSON object whose keys are exactly the language codes and values are the alert texts."
    )
    d, model = generate_json_text(prompt, 0.2)
    if isinstance(d, dict) and all(d.get(c) for c in langs):
        out = {"texts": {c: d[c] for c in langs}, "source": f"gemini:{model}"}
        _ALERT_CACHE[ckey] = out
        return out
    fb = _fallback_alert(phc, summary, items)
    return {"texts": {c: fb[c] for c in langs if c in ("hi", "en")}, "source": "template",
            "missing": [c for c in langs if c not in ("hi", "en")]}

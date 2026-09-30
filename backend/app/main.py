import json
import time
import uuid
from datetime import date
from pathlib import Path
from typing import List, Optional

from fastapi import Depends, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sqlmodel import Session, select

from . import analysis, chat, config, preprocess, stock
from .db import PHC, DvdmsStock, Medicine, Submission, get_session, seed
from .gemini_reader import read_register

app = FastAPI(title="Register Lens API")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


@app.on_event("startup")
def _startup():
    seed()


@app.get("/api/health")
def health():
    return {"ok": True, "gemini_key_set": bool(config.GEMINI_API_KEY), "read_models": config.READ_MODELS,
            "pro_model": config.PRO_MODEL}


@app.get("/api/phcs")
def phcs(district: Optional[str] = None, q: Optional[str] = None, limit: int = 400, session: Session = Depends(get_session)):
    """Facilities for pickers. Demo PHCs (with seeded DVDMS data) come first."""
    query = select(PHC)
    if district:
        query = query.where(PHC.district == district)
    rows = session.exec(query).all()
    if q:
        ql = q.lower()
        rows = [r for r in rows if ql in r.name.lower() or ql in (r.village or "").lower()]
    order = {"DH": 0, "CHC": 1, "PHC": 2, "HWC": 3}
    rows.sort(key=lambda r: (not r.demo, order.get(r.kind, 9), r.name))
    return [r.model_dump() for r in rows[:limit]]


@app.get("/api/districts")
def districts_list(session: Session = Depends(get_session)):
    return sorted({d for d in session.exec(select(PHC.district)).all()})


@app.get("/api/samples")
def samples():
    out = []
    stored = {f.stem for f in config.READINGS_DIR.glob("*.json")}
    for p in sorted(config.SAMPLES_DIR.glob("*.png")) + sorted(config.SAMPLES_DIR.glob("*.jpg")):
        if stored and p.stem not in stored:
            continue  # gallery lists only pages with a stored reading (instant, no API call)
        out.append({"name": p.name, "url": f"/api/sample-image/{p.name}", "thumb": f"/api/sample-thumb/{p.name}"})
    return out


@app.get("/api/sample-image/{name}")
def sample_image(name: str):
    p = (config.SAMPLES_DIR / Path(name).name)
    if not p.exists():
        raise HTTPException(404)
    return FileResponse(p)


@app.get("/api/sample-thumb/{name}")
def sample_thumb(name: str):
    from PIL import Image
    src = config.SAMPLES_DIR / Path(name).name
    if not src.exists():
        raise HTTPException(404)
    dst = config.CACHE_DIR / f"thumb_{src.stem}.jpg"
    if not dst.exists():
        im = Image.open(src).convert("RGB")
        w, h = im.size
        im = im.crop((int(w * 0.08), int(h * 0.08), int(w * 0.92), int(h * 0.92)))
        im.thumbnail((420, 560))
        im.save(dst, quality=80)
    return FileResponse(dst)


@app.get("/api/uploads/{name}")
def upload_file(name: str):
    p = config.UPLOAD_DIR / Path(name).name
    if not p.exists():
        raise HTTPException(404)
    return FileResponse(p)


def _run_scan(image_bytes: bytes, phc_id: str, session: Session, use_cache: bool, tag: str = "", store_path=None, force_phc: Optional[str] = None):
    if not session.get(PHC, phc_id):
        raise HTTPException(404, "Unknown PHC")
    try:
        cleaned, original, steps = preprocess.clean(image_bytes)
    except ValueError as e:
        raise HTTPException(400, str(e))
    sid = uuid.uuid4().hex[:10]
    (config.UPLOAD_DIR / f"{sid}_original.png").write_bytes(original)
    (config.UPLOAD_DIR / f"{sid}_clean.png").write_bytes(cleaned)
    t0 = time.time()
    try:
        page, meta = read_register(cleaned, use_cache=use_cache, tag=tag, store_path=store_path)
    except RuntimeError as e:
        if "RESOURCE_EXHAUSTED" in str(e) or "429" in str(e):
            raise HTTPException(429, "The AI reader has reached its free usage limit for now. Sample pages still work, and uploads will work again after the limit resets.")
        raise HTTPException(502, "The AI reader is busy right now. Please try again in a minute, or try a sample page.")
    if not page.rows:
        raise HTTPException(422, "I could not find a register table in this photo. Try a flatter, well-lit photo of one page.")
    meds = session.exec(select(Medicine)).all()
    used_phc = force_phc or analysis.match_phc(session, page.facility_written) or phc_id
    rows = analysis.prepare_rows([r.model_dump() for r in page.rows], meds, analysis.dvdms_batches(session, used_phc))
    return {
        "phc_id": used_phc,
        "phc_switched": used_phc != phc_id,
        "scan_id": sid,
        "original_url": f"/api/uploads/{sid}_original.png",
        "clean_url": f"/api/uploads/{sid}_clean.png",
        "steps": steps,
        "page_date": page.page_date,
        "facility_written": page.facility_written,
        "rows": rows,
        "meta": {**meta, "total_seconds": round(time.time() - t0, 1)},
        "low_conf_threshold": config.LOW_CONF,
    }


@app.post("/api/scan")
async def scan(file: UploadFile = File(...), phc_id: str = Form(...), session: Session = Depends(get_session)):
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(413, "Image too large (max 15 MB)")
    return _run_scan(data, phc_id, session, use_cache=True)


@app.post("/api/scan-sample")
def scan_sample(name: str = Form(...), phc_id: str = Form(...), session: Session = Depends(get_session)):
    p = config.SAMPLES_DIR / Path(name).name
    if not p.exists():
        raise HTTPException(404, "Sample not found")
    key_file = config.ROOT / "samples" / "answers" / f"{p.stem}.json"
    own_phc = json.loads(key_file.read_text(encoding="utf-8")).get("phc_id") if key_file.exists() else None
    return _run_scan(p.read_bytes(), phc_id, session, use_cache=True,
                     store_path=config.READINGS_DIR / f"{p.stem}.json", force_phc=own_phc)


class ReconcileIn(BaseModel):
    phc_id: str
    rows: List[dict]


@app.post("/api/reconcile")
def reconcile(body: ReconcileIn, session: Session = Depends(get_session)):
    if not session.get(PHC, body.phc_id):
        raise HTTPException(404, "Unknown PHC")
    rec = analysis.reconcile(session, body.phc_id, body.rows)
    stock.save_register_truth(session, body.phc_id, body.rows)
    session.add(Submission(phc_id=body.phc_id, created=date.today().isoformat(), summary=json.dumps(rec["summary"])))
    session.commit()
    return rec


class ExportIn(BaseModel):
    phc_id: str
    rows: List[dict]
    entry_date: Optional[str] = None


@app.post("/api/export")
def export(body: ExportIn, session: Session = Depends(get_session)):
    phc = session.get(PHC, body.phc_id)
    if not phc:
        raise HTTPException(404, "Unknown PHC")
    data = analysis.build_bulk_entry(phc, body.rows, body.entry_date or date.today().isoformat())
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="dvdms_bulk_entry_{phc.id}.xlsx"'})


class AlertIn(BaseModel):
    reconciliation: dict
    languages: Optional[List[str]] = None


@app.post("/api/alert")
def alert(body: AlertIn):
    return analysis.write_alert(body.reconciliation, body.languages)


@app.get("/api/languages")
def languages():
    return {"languages": [{"code": k, "name": v[0], "native": v[1]} for k, v in analysis.LANGUAGES.items()],
            "state_defaults": analysis.STATE_LANGS}


@app.get("/api/up/overview")
def up_overview():
    return stock.state_overview()


@app.get("/api/up/districts.geojson")
def up_geojson():
    f = config.DATA_DIR / "up_districts.geojson"
    if not f.exists():
        raise HTTPException(404)
    return FileResponse(f, media_type="application/geo+json")


@app.get("/api/up/district/{name}")
def up_district(name: str):
    fac = stock.district_facilities(name)
    if not fac:
        raise HTTPException(404, "Unknown district")
    a = _access()
    acc = a.get("districts", {}).get(name)
    catch = a.get("catchment", {})
    fac = [{**f, "villages_served": catch.get(f["id"])} for f in fac]
    return {"district": name, "facilities": fac, "access": acc,
            "note": "Facility locations: OpenStreetMap. Stock is simulated for facilities that have not scanned a register."}


@app.get("/api/facility/{fid}")
def facility(fid: str):
    d = stock.facility_detail(fid)
    if not d:
        raise HTTPException(404)
    return d


_ACCESS: dict = {}


def _access() -> dict:
    if not _ACCESS:
        f = config.DATA_DIR / "up_access.json"
        if f.exists():
            _ACCESS.update(json.loads(f.read_text(encoding="utf-8")))
    return _ACCESS


@app.get("/api/redistribution")
def redistribution(district: str):
    return stock.redistribution(district)


class DecisionIn(BaseModel):
    transfer_id: str
    decision: str
    district: Optional[str] = None


@app.post("/api/redistribution/decision")
def redistribution_decision(body: DecisionIn):
    if body.decision not in ("approved", "rejected", "pending"):
        raise HTTPException(400, "decision must be approved, rejected or pending")
    stock.decide(body.transfer_id, body.decision)
    return {"ok": True}


@app.get("/api/redistribution/orders")
def redistribution_orders(district: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    r = stock.redistribution(district)
    wb = Workbook()
    ws = wb.active
    ws.title = "Transfer orders"
    ws.append(["Order", "Medicine", "Quantity", "Unit", "Batch", "Expiry", "From facility", "To facility", "Distance (km)",
               "Indicative value (INR)", "Stock-out days avoided", "Status"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
    for t in r["transfers"]:
        if t["decision"] != "approved":
            continue
        ws.append([t["id"], t["med_name"], t["qty"], t["unit"], t["batch"], t["expiry"], t["from"]["name"], t["to"]["name"],
                   t["km"], t["value"], t["stockout_days_avoided"], "Approved by district officer"])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col) + 2))
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="transfer_orders_{district}.xlsx"'})


class ChatIn(BaseModel):
    messages: List[dict]
    context: Optional[dict] = None


@app.post("/api/chat")
def chat_endpoint(body: ChatIn):
    return chat.reply(body.messages, body.context)


# Serve the built frontend if present
DIST = config.ROOT / "frontend" / "dist"
if DIST.exists():
    app.mount("/assets", StaticFiles(directory=DIST / "assets"), name="assets")

    @app.get("/{full_path:path}")
    def spa(full_path: str):
        f = DIST / full_path
        if full_path and f.is_file():
            return FileResponse(f)
        return FileResponse(DIST / "index.html")
